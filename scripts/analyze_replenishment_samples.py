from __future__ import annotations

import argparse
import hashlib
import io
import json
import subprocess
import tempfile
import zipfile
import xml.etree.ElementTree as ET
from datetime import date, datetime
from pathlib import Path
from typing import Any

import openpyxl
import xlrd
from pypdf import PdfReader


def scalar(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, bytes):
        return value.hex()
    return value


def sha256_bytes(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()


def extract_xls(content: bytes) -> dict[str, Any]:
    workbook = xlrd.open_workbook(file_contents=content, formatting_info=False)
    sheets: list[dict[str, Any]] = []
    for sheet in workbook.sheets():
        cells = []
        for row_index in range(sheet.nrows):
            for column_index in range(sheet.ncols):
                cell = sheet.cell(row_index, column_index)
                if cell.ctype == xlrd.XL_CELL_EMPTY or cell.value == "":
                    continue
                value: Any = cell.value
                if cell.ctype == xlrd.XL_CELL_DATE:
                    value = xlrd.xldate.xldate_as_datetime(value, workbook.datemode)
                cells.append(
                    {
                        "row": row_index + 1,
                        "column": column_index + 1,
                        "address": f"{xlrd.formula.colname(column_index)}{row_index + 1}",
                        "value": scalar(value),
                        "cellType": int(cell.ctype),
                    }
                )
        sheets.append(
            {
                "name": sheet.name,
                "visibility": int(sheet.visibility),
                "rowCount": sheet.nrows,
                "columnCount": sheet.ncols,
                "mergedRanges": [
                    {
                        "rowStart": row_start + 1,
                        "rowEndExclusive": row_end + 1,
                        "columnStart": column_start + 1,
                        "columnEndExclusive": column_end + 1,
                    }
                    for row_start, row_end, column_start, column_end in sheet.merged_cells
                ],
                "cells": cells,
            }
        )
    return {"format": "xls", "sheets": sheets}


def extract_xlsx(content: bytes) -> dict[str, Any]:
    formulas = openpyxl.load_workbook(io.BytesIO(content), data_only=False, read_only=False)
    values = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=False)
    sheets: list[dict[str, Any]] = []
    for formula_sheet in formulas.worksheets:
        value_sheet = values[formula_sheet.title]
        cells = []
        for row in formula_sheet.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                cached_value = value_sheet[cell.coordinate].value
                cells.append(
                    {
                        "row": cell.row,
                        "column": cell.column,
                        "address": cell.coordinate,
                        "value": scalar(cell.value),
                        "cachedValue": scalar(cached_value),
                        "dataType": cell.data_type,
                        "numberFormat": cell.number_format,
                    }
                )
        sheets.append(
            {
                "name": formula_sheet.title,
                "state": formula_sheet.sheet_state,
                "rowCount": formula_sheet.max_row,
                "columnCount": formula_sheet.max_column,
                "mergedRanges": [str(item) for item in formula_sheet.merged_cells.ranges],
                "cells": cells,
            }
        )
    return {"format": "xlsx", "sheets": sheets}


def extract_spreadsheetml(content: bytes) -> dict[str, Any]:
    root = ET.fromstring(content)
    uri = "urn:schemas-microsoft-com:office:spreadsheet"
    q = lambda name: f"{{{uri}}}{name}"
    sheets = []
    for worksheet in root.findall(q("Worksheet")):
        cells, row_number, max_column = [], 0, 0
        table = worksheet.find(q("Table"))
        for row in [] if table is None else table.findall(q("Row")):
            row_number = int(row.get(q("Index"), row_number + 1))
            column_number = 0
            for cell in row.findall(q("Cell")):
                column_number = int(cell.get(q("Index"), column_number + 1))
                data = cell.find(q("Data"))
                value = "" if data is None or data.text is None else data.text
                formula = cell.get(q("Formula"))
                if formula is not None or value != "":
                    cells.append({"row": row_number, "column": column_number,
                        "address": f"{xlrd.formula.colname(column_number - 1)}{row_number}",
                        "value": value, "formula": formula,
                        "dataType": None if data is None else data.get(q("Type"))})
                max_column = max(max_column, column_number)
                column_number += int(cell.get(q("MergeAcross"), "0"))
        sheets.append({"name": worksheet.get(q("Name")), "rowCount": row_number,
            "columnCount": max_column, "cells": cells})
    return {"format": "spreadsheetml-2003", "sheets": sheets}


def extract_pdf(content: bytes) -> dict[str, Any]:
    reader = PdfReader(io.BytesIO(content))
    return {
        "format": "pdf",
        "pageCount": len(reader.pages),
        "pages": [
            {"page": index + 1, "text": page.extract_text() or ""}
            for index, page in enumerate(reader.pages)
        ],
    }


def extract_image(path: Path) -> dict[str, Any]:
    attempts = []
    for language in ("chi_sim+eng", "eng"):
        result = subprocess.run(
            ["tesseract", str(path), "stdout", "-l", language],
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            check=False,
        )
        attempts.append(
            {
                "language": language,
                "exitCode": result.returncode,
                "text": result.stdout,
                "error": result.stderr,
            }
        )
        if result.returncode == 0:
            break
    return {"format": "image", "ocrAttempts": attempts}


def extract_content(name: str, content: bytes, source_path: Path | None = None) -> dict[str, Any]:
    suffix = Path(name).suffix.lower()
    if suffix in {".xls", ".xlsx"} and content.startswith(b"PK"):
        return extract_xlsx(content)
    if suffix in {".xls", ".xlsx"} and content.lstrip().startswith(b"<?xml"):
        return extract_spreadsheetml(content)
    if suffix == ".xls":
        return extract_xls(content)
    if suffix == ".xlsx":
        return extract_xlsx(content)
    if suffix == ".pdf":
        return extract_pdf(content)
    if suffix in {".png", ".jpg", ".jpeg", ".tif", ".tiff"}:
        if source_path is not None:
            return extract_image(source_path)
        with tempfile.TemporaryDirectory() as temp_directory:
            temporary_path = Path(temp_directory) / Path(name).name
            temporary_path.write_bytes(content)
            return extract_image(temporary_path)
    return {"format": suffix.removeprefix(".") or "unknown", "status": "not_parsed"}


def extract_zip(content: bytes) -> dict[str, Any]:
    entries = []
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        for info in archive.infolist():
            entry_content = archive.read(info.filename)
            entries.append(
                {
                    "name": info.filename,
                    "sizeBytes": info.file_size,
                    "sha256": sha256_bytes(entry_content),
                    "content": extract_content(info.filename, entry_content),
                }
            )
    return {"format": "zip", "entries": entries}


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()

    files = []
    for path in sorted(item for item in args.source.iterdir() if item.is_file()):
        content = path.read_bytes()
        try:
            extracted = (
                extract_zip(content)
                if path.suffix.lower() == ".zip"
                else extract_content(path.name, content, path)
            )
            parse_error = None
        except Exception as error:  # Preserve failures in the evidence manifest.
            extracted = None
            parse_error = f"{type(error).__name__}: {error}"
        files.append(
            {
                "name": path.name,
                "sizeBytes": len(content),
                "sha256": sha256_bytes(content),
                "lastModified": datetime.fromtimestamp(path.stat().st_mtime).astimezone().isoformat(),
                "extracted": extracted,
                "parseError": parse_error,
            }
        )

    result = {
        "sourceDirectory": str(args.source),
        "generatedAt": datetime.now().astimezone().isoformat(),
        "fileCount": len(files),
        "files": files,
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"fileCount": len(files), "output": str(args.output)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
